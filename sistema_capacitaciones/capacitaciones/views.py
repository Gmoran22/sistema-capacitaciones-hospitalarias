import os
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render, redirect
from django.contrib import messages
from django.db import transaction
from django.core.mail import EmailMessage
from django.conf import settings
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required

from .models import Capacitado
from .forms import CapacitadoForm

logger = logging.getLogger(__name__)

def enviar_material_capacitacion(destinatario, nombre_completo):
    asunto = "Material de Capacitación y Alta de Sistema"
    cuerpo = (
        f"Hola {nombre_completo},\n\n"
        "Confirmamos que tu registro de capacitación fue recibido y procesado correctamente.\n"
        "Adjunto a este correo encontrarás el manual instructivo de uso del sistema.\n\n"
        "Saludos cordiales,\n"
        "Equipo de Implementación y Capacitación de Sistemas Sanitarios."
    )
    
    ruta_pdf = os.path.join(settings.BASE_DIR, 'manual_capacitacion.pdf')

    try:
        if settings.EMAIL_HOST_USER:
            email = EmailMessage(
                subject=asunto,
                body=cuerpo,
                from_email=settings.DEFAULT_FROM_EMAIL,
                to=[destinatario]
            )
            if os.path.exists(ruta_pdf):
                email.attach_file(ruta_pdf)

            email.send(fail_silently=True)
            return True
        return False
    except Exception as e:
        logger.error(f"Falla al despachar email a {destinatario}: {str(e)}")
        return False


def registro_capacitado_view(request):
    if request.method == 'POST':
        form = CapacitadoForm(request.POST)
        if form.is_valid():
            try:
                with transaction.atomic():
                    capacitado = form.save(commit=False)
                    nombre_completo = f"{capacitado.nombre} {capacitado.apellido}"
                    envio_ok = enviar_material_capacitacion(capacitado.mail_personal, nombre_completo)
                    capacitado.material_enviado = envio_ok
                    capacitado.save()

                return redirect('registro_exitoso')

            except Exception as e:
                logger.error(f"Error al guardar registro: {str(e)}")
                messages.error(request, "Ocurrió un error al procesar el registro. Inténtelo nuevamente.")
    else:
        form = CapacitadoForm()

    return render(request, 'capacitaciones/registro.html', {'form': form})


def registro_exitoso_view(request):
    return render(request, 'capacitaciones/exito.html')


@login_required
def exportar_excel_view(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nómina Capacitaciones"

    headers = [
        "Tipo de Usuario", "¿Acceso SIGEHOS?", "Trabaja en otro hospital? (Cuál?)",
        "Permisos Solicitados", "Primer Apellido", "Segundo Apellido", "Primer Nombre", "Segundo Nombre",
        "CUIL", "Correo GCBA", "Mail Personal", "Teléfono", "Profesión", "Categoría",
        "Tipo Matrícula", "N° Matrícula", "Especialidad", "Servicio", "Área / Sala", "Días y Horarios",
        "Fecha Registro", "Material Enviado"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    registros = Capacitado.objects.all().order_by('fecha_registro')
    for reg in registros:
        ws.append([
            reg.get_tipo_usuario_display(),
            "SÍ" if reg.tiene_acceso_sigehos else "NO",
            reg.trabaja_otro_hospital,
            reg.get_permisos_solicitar_display(),
            reg.apellido,
            reg.otro_apellido or "-",
            reg.nombre,
            reg.otro_nombre or "-",
            reg.cuil,
            reg.correo_gcba or "-",
            reg.mail_personal,
            reg.telefono or "-",
            reg.profesion,
            reg.get_categoria_display(),
            reg.get_tipo_matricula_display(),
            reg.numero_matricula,
            reg.especialidad or "-",
            reg.servicio,
            reg.area_trabajo,
            reg.dias_horarios,
            reg.fecha_registro.strftime("%d/%m/%Y %H:%M"),
            "SÍ" if reg.material_enviado else "NO"
        ])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Nomina_Oficial_Capacitados.xlsx"'
    wb.save(response)
    return response